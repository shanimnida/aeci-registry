"""Admin-level (upload_view / download_template_view) tests for the
spreadsheet import path -- everything test_spreadsheet_parsing.py checks at
the function level, checked again through the real upload screen and
permission checks, plus the end-to-end "download the blank template, fill
it in, upload it, approve it" path the volunteer actually follows.
"""

import io
from pathlib import Path

import openpyxl
import pytest
from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from committees.models import CommitteeMembership
from core import groups
from imports.models import ImportBatch, StagedPerson, StagedPersonStatus
from imports.spreadsheet import HEADER_TEXTS, write_blank_template_bytes
from people.models import Person

from .spreadsheet_helpers import csv_bytes, person_row, xlsx_bytes

FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture
def secretariat_user(db):
    user = User.objects.create_user("sec", password="x", is_staff=True)
    user.groups.add(Group.objects.get(name=groups.SECRETARIAT))
    return user


@pytest.fixture
def chairperson_user(db):
    user = User.objects.create_user("chair", password="x", is_staff=True)
    user.groups.add(Group.objects.get(name=groups.CHAIRPERSON))
    return user


def _upload(name, raw, content_type):
    return SimpleUploadedFile(name, raw, content_type=content_type)


@pytest.mark.django_db
def test_uploading_a_csv_file_stages_the_right_number_of_people(client, secretariat_user):
    client.force_login(secretariat_user)
    rows = [
        person_row(**{"Source Image": "IMG_A.jpg", "Last Name": "CRUZ", "First Name": "ANA"}),
        person_row(**{"Source Image": "IMG_B.jpg", "Last Name": "REYES", "First Name": "BEN"}),
    ]
    response = client.post(
        reverse("admin:imports_importbatch_upload"),
        {"file": _upload("profiling.csv", csv_bytes(rows), "text/csv")},
    )
    assert response.status_code == 302
    assert ImportBatch.objects.count() == 1
    batch = ImportBatch.objects.get()
    assert batch.rows.count() == 2
    assert Person.objects.count() == 0


@pytest.mark.django_db
def test_uploading_an_xlsx_file_stages_the_right_number_of_people(client, secretariat_user):
    client.force_login(secretariat_user)
    rows = [
        person_row(**{"Source Image": "IMG_A.jpg", "Last Name": "CRUZ", "First Name": "ANA"}),
        person_row(**{"Source Image": "IMG_B.jpg", "Last Name": "REYES", "First Name": "BEN"}),
    ]
    response = client.post(
        reverse("admin:imports_importbatch_upload"),
        {
            "file": _upload(
                "profiling.xlsx",
                xlsx_bytes(rows),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 302
    assert ImportBatch.objects.count() == 1
    assert ImportBatch.objects.get().rows.count() == 2


@pytest.mark.django_db
def test_xlsx_and_csv_uploads_of_the_same_content_stage_identical_raw_data(
    client, secretariat_user
):
    rows = [
        person_row(
            **{
                "Music and Arts": "X",
                "Child 1 Name": "JUANA SANTOS",
                "Child 1 Date of Birth": "2015-05-05",
            }
        )
    ]
    client.force_login(secretariat_user)
    client.post(
        reverse("admin:imports_importbatch_upload"),
        {"file": _upload("a.csv", csv_bytes(rows), "text/csv")},
    )
    client.post(
        reverse("admin:imports_importbatch_upload"),
        {
            "file": _upload(
                "b.xlsx",
                xlsx_bytes(rows),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert ImportBatch.objects.count() == 2
    raw_data_sets = [batch.rows.get().raw_data for batch in ImportBatch.objects.all()]
    assert raw_data_sets[0] == raw_data_sets[1]


@pytest.mark.django_db
def test_a_spreadsheet_with_wrong_columns_is_rejected_and_stages_nothing(
    client, secretariat_user
):
    client.force_login(secretariat_user)
    bad_csv = b"Source Image,Last Name,First Name\r\nIMG_1.jpg,CRUZ,ANA\r\n"
    response = client.post(
        reverse("admin:imports_importbatch_upload"),
        {"file": _upload("wrong.csv", bad_csv, "text/csv")},
    )
    assert response.status_code == 200  # re-rendered with errors, no redirect
    body = response.content.decode()
    assert "missing column" in body
    assert ImportBatch.objects.count() == 0
    assert StagedPerson.objects.count() == 0


@pytest.mark.django_db
def test_a_spreadsheet_row_error_names_the_row_and_column_on_the_upload_screen(
    client, secretariat_user
):
    client.force_login(secretariat_user)
    rows = [
        person_row(**{"Source Image": "IMG_A.jpg"}),
        person_row(**{"Source Image": "IMG_B.jpg", "First Name": ""}),  # row 3: bad
    ]
    response = client.post(
        reverse("admin:imports_importbatch_upload"),
        {"file": _upload("profiling.csv", csv_bytes(rows), "text/csv")},
    )
    assert response.status_code == 200
    body = response.content.decode()
    assert "row 3" in body
    assert "First Name" in body
    assert ImportBatch.objects.count() == 0


@pytest.mark.django_db
def test_the_blank_template_download_serves_an_xlsx_with_the_documented_headers(
    client, secretariat_user
):
    client.force_login(secretariat_user)
    response = client.get(reverse("admin:imports_importbatch_download_template"))
    assert response.status_code == 200
    assert response["Content-Type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in response["Content-Disposition"]
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    header_row = next(workbook.active.iter_rows(min_row=1, max_row=1, values_only=True))
    assert list(header_row) == list(HEADER_TEXTS)


@pytest.mark.django_db
def test_a_chairperson_cannot_download_the_template(client, chairperson_user):
    client.force_login(chairperson_user)
    response = client.get(reverse("admin:imports_importbatch_download_template"))
    assert response.status_code == 403


@pytest.mark.django_db
def test_an_anonymous_user_cannot_download_the_template(client):
    response = client.get(reverse("admin:imports_importbatch_download_template"))
    assert response.status_code == 302
    assert "login" in response.url


@pytest.mark.django_db
def test_downloading_the_template_filling_it_in_and_approving_creates_a_real_person(
    client, secretariat_user
):
    """End to end: exactly the volunteer's workflow -- download the blank
    template, fill in one fictional person by hand (as if in Excel), upload
    it, then approve the staged row on the review screen. Proves the
    template's own headers and the parser agree, and that a spreadsheet-
    staged row survives the full approval path (full_clean() and all)
    exactly like a JSON-staged one does.
    """
    client.force_login(secretariat_user)

    template_response = client.get(reverse("admin:imports_importbatch_download_template"))
    workbook = openpyxl.load_workbook(io.BytesIO(template_response.content))
    sheet = workbook.active
    filled = person_row(
        **{
            "Source Image": "IMG_END_TO_END.jpg",
            "Last Name": "VILLANUEVA",
            "First Name": "ROSA",
            "Gender": "FEMALE",
            "Civil Status": "SINGLE",
            "Music and Arts": "X",
            "Certification Date": "2026-08-16",
        }
    )
    sheet.append([filled[header] for header in HEADER_TEXTS])
    buffer = io.BytesIO()
    workbook.save(buffer)

    upload_response = client.post(
        reverse("admin:imports_importbatch_upload"),
        {
            "file": _upload(
                "filled_template.xlsx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert upload_response.status_code == 302
    batch = ImportBatch.objects.get()
    row = batch.rows.get()
    assert row.raw_data["last_name"] == "VILLANUEVA"

    review_url = reverse("admin:imports_importbatch_review", args=[batch.pk, row.pk])
    data = row.raw_data
    post = {
        "action": "approve",
        "member_no": "",
        "last_name": data["last_name"],
        "first_name": data["first_name"],
        "middle_name": "",
        "suffix": "",
        "date_of_birth": "",
        "place_of_birth": "",
        "gender": data["gender"],
        "civil_status": data["civil_status"],
        "nationality": "",
        "home_address": "",
        "mobile_number": "",
        "email": "",
        "spouse_name": "",
        "date_of_marriage": "",
        "form_version": data["form_version"],
        "date_filed": "",
        "certification_date": data["certification_date"],
        "emergency_contact_name": "",
        "emergency_relationship": "",
        "emergency_number": "",
        "committees": data["committees"],
        "membership_status": "MEMBER",
        "children-TOTAL_FORMS": "0",
        "children-INITIAL_FORMS": "0",
        "children-MIN_NUM_FORMS": "0",
        "children-MAX_NUM_FORMS": "1000",
    }
    approve_response = client.post(review_url, post)
    assert approve_response.status_code == 302

    person = Person.objects.get(last_name="Villanueva", first_name="Rosa")
    assert person.gender == "FEMALE"
    assert CommitteeMembership.objects.filter(person=person, committee__name="Music and Arts").exists()
    row.refresh_from_db()
    assert row.status == StagedPersonStatus.APPROVED
