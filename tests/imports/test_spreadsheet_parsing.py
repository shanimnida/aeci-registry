"""Unit tests for imports/spreadsheet.py: parsing .xlsx and .csv uploads
into the same entry shape imports/parsing.parse_import_json returns, with
errors that name the spreadsheet row and the column header rather than a
JSON key. All fixtures are built in-memory (see spreadsheet_helpers.py) and
every person in them is fictional -- never the real data under
`form images/`.
"""

import openpyxl
import pytest

from imports.parsing import ImportValidationError
from imports.spreadsheet import (
    HEADER_TEXTS,
    parse_import_file,
    sniff_import_format,
    write_blank_template_bytes,
)

from .spreadsheet_helpers import csv_bytes, person_row, row_dict, xlsx_bytes


# -- format sniffing ----------------------------------------------------------

def test_sniff_recognises_xlsx_by_its_zip_magic_bytes():
    assert sniff_import_format(xlsx_bytes([person_row()])) == "xlsx"


def test_sniff_recognises_csv_as_the_fallback():
    assert sniff_import_format(csv_bytes([person_row()])) == "csv"


def test_sniff_recognises_json_by_its_leading_bracket():
    assert sniff_import_format(b'[{"source_image": "x.jpg"}]') == "json"
    assert sniff_import_format(b'  [{"a": 1}]') == "json"


# -- xlsx and csv parity -------------------------------------------------------

def test_xlsx_and_csv_with_the_same_content_parse_into_identical_entries():
    rows = [
        person_row(
            **{
                "Music and Arts": "X",
                "Events": "x",  # lower-case tick still counts
                "Child 1 Name": "JUANA SANTOS",
                "Child 1 Date of Birth": "2015-05-05",
            }
        )
    ]
    csv_entries = parse_import_file(csv_bytes(rows))
    xlsx_entries = parse_import_file(xlsx_bytes(rows))
    assert csv_entries == xlsx_entries
    assert csv_entries[0]["committees"] == ["Music and Arts", "Events"]
    assert csv_entries[0]["children"] == [
        {"full_name": "JUANA SANTOS", "date_of_birth": "2015-05-05"}
    ]


def test_a_well_formed_spreadsheet_parses_into_the_right_number_of_entries():
    rows = [
        person_row(**{"Source Image": "IMG_A.jpg", "Last Name": "CRUZ", "First Name": "ANA"}),
        person_row(**{"Source Image": "IMG_B.jpg", "Last Name": "REYES", "First Name": "BEN"}),
        person_row(**{"Source Image": "IMG_C.jpg", "Last Name": "TORRES", "First Name": "CY"}),
    ]
    entries = parse_import_file(csv_bytes(rows))
    assert len(entries) == 3
    assert entries[0]["last_name"] == "CRUZ"
    assert entries[2]["last_name"] == "TORRES"


def test_a_wholly_blank_trailing_row_is_skipped_not_staged_as_an_entry():
    rows = [person_row(), row_dict()]  # second row: every column blank
    entries = parse_import_file(csv_bytes(rows))
    assert len(entries) == 1


# -- row/column error reporting ------------------------------------------------

def test_an_error_names_the_spreadsheet_row_and_the_column_header():
    rows = [
        person_row(**{"Source Image": "IMG_A.jpg"}),
        person_row(**{"Source Image": "IMG_B.jpg", "Last Name": ""}),  # bad: row 3
        person_row(**{"Source Image": "IMG_C.jpg"}),
    ]
    with pytest.raises(ImportValidationError) as exc_info:
        parse_import_file(csv_bytes(rows))
    errors = exc_info.value.errors
    assert len(errors) == 1
    assert "row 3" in errors[0]
    assert "IMG_B.jpg" in errors[0]
    assert "'Last Name'" in errors[0]
    # The internal JSON field name must not leak into a spreadsheet error.
    assert "last_name" not in errors[0]


def test_a_bad_date_names_its_own_column_header_not_a_json_key():
    rows = [person_row(**{"Date of Birth": "not-a-date"})]
    with pytest.raises(ImportValidationError) as exc_info:
        parse_import_file(csv_bytes(rows))
    assert "'Date of Birth'" in exc_info.value.errors[0]


def test_a_bad_child_slot_names_that_slots_own_headers():
    """A date with no name in the same slot -- 'full_name is required' must
    say 'Child 2 Name', not the JSON key 'full_name'."""
    rows = [person_row(**{"Child 2 Date of Birth": "2020-01-01"})]
    with pytest.raises(ImportValidationError) as exc_info:
        parse_import_file(csv_bytes(rows))
    errors = exc_info.value.errors
    assert any("Child 2 Name" in e for e in errors)


def test_every_problem_in_the_file_is_reported_together():
    rows = [
        person_row(**{"Last Name": ""}),
        person_row(**{"Confidence": "very high"}),
    ]
    with pytest.raises(ImportValidationError) as exc_info:
        parse_import_file(csv_bytes(rows))
    assert len(exc_info.value.errors) == 2


def test_an_empty_spreadsheet_after_the_header_row_is_rejected():
    with pytest.raises(ImportValidationError) as exc_info:
        parse_import_file(csv_bytes([]))
    assert "no entries" in exc_info.value.errors[0]


# -- wrong columns --------------------------------------------------------------

def test_a_file_missing_expected_columns_is_refused_clearly():
    import csv as _csv
    import io as _io

    buffer = _io.StringIO()
    writer = _csv.writer(buffer)
    writer.writerow(["Source Image", "Last Name", "First Name"])
    writer.writerow(["IMG_1.jpg", "CRUZ", "ANA"])
    with pytest.raises(ImportValidationError) as exc_info:
        parse_import_file(buffer.getvalue().encode("utf-8"))
    assert "missing column" in exc_info.value.errors[0]
    assert "Confidence" in exc_info.value.errors[0]


def test_a_file_with_an_unrecognised_column_is_refused_clearly():
    import csv as _csv
    import io as _io

    buffer = _io.StringIO()
    writer = _csv.writer(buffer)
    headers = list(HEADER_TEXTS)
    headers[headers.index("First Name")] = "First Nam"  # typo
    writer.writerow(headers)
    writer.writerow(["x"] * len(headers))
    with pytest.raises(ImportValidationError) as exc_info:
        parse_import_file(buffer.getvalue().encode("utf-8"))
    joined = "\n".join(exc_info.value.errors)
    assert "does not recognise" in joined
    assert "First Nam" in joined


# -- Grievance and Reconciliation (Board-appointed, never on this form) -------

def test_grievance_and_reconciliation_ticked_is_recorded_and_flagged_not_rejected():
    rows = [person_row(**{"Grievance and Reconciliation": "X"})]
    entries = parse_import_file(csv_bytes(rows))
    assert entries[0]["committees"] == ["Grievance and Reconciliation"]
    flags = entries[0]["uncertain_fields"]
    assert any(
        item["field"] == "committees" and "Board-appointed" in item["why"] for item in flags
    )


def test_grievance_ticked_alongside_other_committees_still_lists_all_of_them():
    rows = [person_row(**{"Music and Arts": "X", "Grievance and Reconciliation": "X"})]
    entries = parse_import_file(csv_bytes(rows))
    assert entries[0]["committees"] == ["Music and Arts", "Grievance and Reconciliation"]


# -- Flagged Fields column -> uncertain_fields ---------------------------------

def test_flagged_fields_column_maps_a_recognised_header_to_its_field():
    rows = [
        person_row(
            **{"Flagged Fields": "Date of Birth: read as smudged, year unclear"}
        )
    ]
    entries = parse_import_file(csv_bytes(rows))
    flags = entries[0]["uncertain_fields"]
    assert flags == [{"field": "date_of_birth", "why": "read as smudged, year unclear"}]


def test_flagged_fields_column_supports_multiple_entries_separated_by_pipe():
    rows = [
        person_row(
            **{
                "Flagged Fields": (
                    "Date of Birth: year digit unclear | Committees: three boxes ticked"
                )
            }
        )
    ]
    entries = parse_import_file(csv_bytes(rows))
    fields = {item["field"] for item in entries[0]["uncertain_fields"]}
    assert fields == {"date_of_birth", "committees"}


def test_a_flagged_fields_entry_with_no_recognised_header_is_kept_as_a_note():
    rows = [person_row(**{"Flagged Fields": "handwriting very faint throughout"})]
    entries = parse_import_file(csv_bytes(rows))
    assert entries[0]["uncertain_fields"] == [
        {"field": "Note", "why": "handwriting very faint throughout"}
    ]


def test_a_child_column_flag_collapses_to_the_children_section_flag():
    rows = [
        person_row(**{"Flagged Fields": "Child 1 Date of Birth: corrected on the page"})
    ]
    entries = parse_import_file(csv_bytes(rows))
    assert entries[0]["uncertain_fields"][0]["field"] == "children"


# -- blank template -------------------------------------------------------------

def test_the_blank_template_has_exactly_the_documented_headers_in_a_bold_row():
    raw = write_blank_template_bytes()
    import io as _io

    workbook = openpyxl.load_workbook(_io.BytesIO(raw))
    sheet = workbook.active
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    assert list(header_row) == list(HEADER_TEXTS)


def test_a_spreadsheet_built_from_the_blank_template_headers_round_trips():
    """The blank template's own header row, filled in with one fictional
    person, must parse cleanly -- proof the template and the parser agree."""
    raw = write_blank_template_bytes()
    import io as _io

    workbook = openpyxl.load_workbook(_io.BytesIO(raw))
    sheet = workbook.active
    row = person_row(**{"Last Name": "VILLANUEVA", "First Name": "ROSA"})
    sheet.append([row[header] for header in HEADER_TEXTS])
    buffer = _io.BytesIO()
    workbook.save(buffer)

    entries = parse_import_file(buffer.getvalue())
    assert len(entries) == 1
    assert entries[0]["last_name"] == "VILLANUEVA"


# -- JSON still works through the same dispatcher ------------------------------

def test_a_json_file_still_parses_through_the_shared_dispatcher():
    raw = (
        b'[{"source_image": "IMG_1.jpg", "form_version": "v1", "last_name": "CRUZ", '
        b'"first_name": "ANA", "confidence": "high", "committees": [], "children": [], '
        b'"uncertain_fields": []}]'
    )
    entries = parse_import_file(raw)
    assert len(entries) == 1
    assert entries[0]["last_name"] == "CRUZ"
