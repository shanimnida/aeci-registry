"""Validates and parses an uploaded .xlsx or .csv spreadsheet against the
column layout docs/IMPORT_TEMPLATE.md documents, and the entry point both
imports/admin.py's upload_view and the tests use to accept a file of any
supported format -- .xlsx, .csv, or the original .json.

Why a spreadsheet at all: the JSON produced by the AI is not something a
non-technical volunteer can review. A spreadsheet is -- they open it in
Excel, scan it against the paper forms in hand, correct what the AI got
wrong, and save it back as .xlsx or .csv. This module's job mirrors
imports/parsing.py's for the JSON path: never touch the database, never
touch people.Person, and if anything about the file will not stage, say
exactly which row and which column so the volunteer can fix the file (or
the AI's transcription) rather than stare at a stack trace.

**The spreadsheet mirrors the paper form on purpose.** One column per
committee (so a volunteer marks a box, never has to spell a committee's
name), five fixed child slots (the paper form has exactly five rows), and
every other field in the order the printed sections lay them out -- so
checking a row on screen against a sheet in hand is a straight scan, not a
hunt. See HEADER_TEXTS below for the exact columns, and
docs/IMPORT_TEMPLATE.md for why each one is named the way it is.

Every row this module builds has *exactly* the same shape
imports/parsing.parse_import_json returns for one JSON entry, and is
validated by the exact same `_validate_entry` -- so imports/services.py and
imports/forms.py need not know or care which file format a batch was
staged from.
"""

import csv
import datetime
import io
import re
import zipfile

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException

from .parsing import (
    ALL_COMMITTEE_NAMES,
    APPOINTED_ONLY_COMMITTEE_NAME,
    ImportValidationError,
    _validate_entry,
    parse_import_json,
)

CHILD_SLOTS = 5

# -- column layout ---------------------------------------------------------
# Single source of truth for: the blank template's header row, the header
# set an uploaded file is checked against, which internal field each column
# feeds, and the column-header text an error names instead of a raw JSON
# key. Order here is display order only -- an uploaded file's columns are
# matched by header *text*, not position, so a volunteer who drags a column
# sideways in Excel is not punished for it.
_PERSONAL_FIELDS = (
    ("Member No", "member_no"),
    ("Last Name", "last_name"),
    ("First Name", "first_name"),
    ("Middle Name", "middle_name"),
    ("Suffix", "suffix"),
    ("Date of Birth", "date_of_birth"),
    ("Place of Birth", "place_of_birth"),
    ("Gender", "gender"),
    ("Civil Status", "civil_status"),
    ("Nationality", "nationality"),
)
_CONTACT_FIELDS = (
    ("Home Address", "home_address"),
    ("Mobile Number", "mobile_number"),
    ("Email", "email"),
)
_FAMILY_FIELDS_BEFORE_CHILDREN = (
    ("Spouse Name", "spouse_name"),
    ("Date of Marriage", "date_of_marriage"),
)
_FAMILY_FIELDS_AFTER_CHILDREN = (
    ("Emergency Contact Name", "emergency_contact_name"),
    ("Emergency Contact Relationship", "emergency_relationship"),
    ("Emergency Contact Number", "emergency_number"),
)
_CERTIFICATION_FIELDS = (
    ("Date Filed", "date_filed"),
    ("Certification Date", "certification_date"),
    ("Form Version", "form_version"),
)

_COLUMNS = []


def _field_column(header, key):
    _COLUMNS.append({"header": header, "kind": "field", "key": key})


def _child_columns(n):
    _COLUMNS.append({"header": f"Child {n} Name", "kind": "child_name", "n": n})
    _COLUMNS.append({"header": f"Child {n} Date of Birth", "kind": "child_dob", "n": n})


def _committee_column(name):
    _COLUMNS.append({"header": name, "kind": "committee", "name": name})


_field_column("Source Image", "source_image")
for _header, _key in _PERSONAL_FIELDS:
    _field_column(_header, _key)
for _header, _key in _CONTACT_FIELDS:
    _field_column(_header, _key)
for _header, _key in _FAMILY_FIELDS_BEFORE_CHILDREN:
    _field_column(_header, _key)
for _n in range(1, CHILD_SLOTS + 1):
    _child_columns(_n)
for _header, _key in _FAMILY_FIELDS_AFTER_CHILDREN:
    _field_column(_header, _key)
for _name in ALL_COMMITTEE_NAMES:
    _committee_column(_name)
for _header, _key in _CERTIFICATION_FIELDS:
    _field_column(_header, _key)
_field_column("Notes", "notes")
_field_column("Confidence", "confidence")
_COLUMNS.append({"header": "Flagged Fields", "kind": "flagged_fields"})

HEADER_TEXTS = tuple(column["header"] for column in _COLUMNS)
_FIELD_COLUMNS = tuple((c["header"], c["key"]) for c in _COLUMNS if c["kind"] == "field")

# field key -> column header, for error messages (imports/parsing.py's
# field_labels) so a validation error names 'Last Name', not 'last_name'.
FIELD_LABELS = {key: header for header, key in _FIELD_COLUMNS}
for _n in range(1, CHILD_SLOTS + 1):
    FIELD_LABELS[f"child_{_n}_name"] = f"Child {_n} Name"
    FIELD_LABELS[f"child_{_n}_date_of_birth"] = f"Child {_n} Date of Birth"

# column header (lowercased) -> the field a Flagged Fields entry naming that
# header should attach to on the review screen (imports/admin.py's
# uncertain_by_field). Child columns and every committee column collapse to
# the one section-level flag review.html already knows how to show
# ("children" / "committees") -- there is no per-child or per-committee flag
# slot in that screen, on purpose (docs/IMPORT_TEMPLATE.md: flags are
# per-section, not per-checkbox).
_FLAG_FIELD_BY_HEADER = {header.lower(): key for header, key in _FIELD_COLUMNS}
for _n in range(1, CHILD_SLOTS + 1):
    _FLAG_FIELD_BY_HEADER[f"child {_n} name".lower()] = "children"
    _FLAG_FIELD_BY_HEADER[f"child {_n} date of birth".lower()] = "children"
for _name in ALL_COMMITTEE_NAMES:
    _FLAG_FIELD_BY_HEADER[_name.lower()] = "committees"
_FLAG_FIELD_BY_HEADER["children"] = "children"
_FLAG_FIELD_BY_HEADER["committees"] = "committees"


# -- cell normalisation ------------------------------------------------------

def _normalize_cell(value):
    """One cell -> text or None, regardless of whether openpyxl handed back
    a string, a number, a bool, or a datetime -- Excel silently retypes a
    cell depending on how it was entered, and that must never read as a
    missing or malformed field. Never invents or corrects *content*: a date
    typed as a real Excel date becomes the same YYYY-MM-DD text the AI would
    have written; nothing here changes what a value means, only how it is
    represented as text.
    """
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, bool):
        # openpyxl reads an Excel checkbox-style TRUE as bool True. Only
        # committee columns are ever ticked this way; True still means
        # "ticked" there.
        return "X" if value else None
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    text = str(value).strip()
    return text or None


_FLAG_SPLIT_RE = re.compile(r"\s*(?:\||\r\n|\r|\n)\s*")


def _parse_flagged_fields(text: str | None) -> list[dict]:
    """'Date of Birth: read as OCT 14 1092, year digit unclear | Committees:
    three boxes ticked' -> the same [{'field': ..., 'why': ...}] shape
    parse_import_json's `uncertain_fields` uses, so imports/admin.py's
    per-field flag display (uncertain_by_field) needs no changes at all to
    work from a spreadsheet. Entries are separated by '|' or a line break --
    Excel cells commonly hold either. A segment with no 'Header: text' shape
    is kept, not dropped, under the field name 'Note' so nothing the AI or
    volunteer wrote silently vanishes from the review screen.
    """
    if not text:
        return []
    segments = [s.strip() for s in _FLAG_SPLIT_RE.split(text) if s.strip()]
    parsed = []
    for segment in segments:
        if ":" in segment:
            header_part, why_part = segment.split(":", 1)
            header_part = header_part.strip()
            why_part = why_part.strip()
            field = _FLAG_FIELD_BY_HEADER.get(header_part.lower(), header_part)
        else:
            field = "Note"
            why_part = segment
        parsed.append({"field": field, "why": why_part})
    return parsed


# -- header (column) validation ---------------------------------------------

def _validate_headers(headers_present: set[str]) -> list[str]:
    expected = set(HEADER_TEXTS)
    missing = [h for h in HEADER_TEXTS if h not in headers_present]
    unexpected = sorted(headers_present - expected)
    errors = []
    if missing:
        errors.append(
            "This file is missing column(s) AEGIS expects: " + ", ".join(missing) + ". "
            "Download the blank import template and copy your data into it, or ask the AI "
            "to use the exact headers in docs/IMPORT_TEMPLATE.md."
        )
    if unexpected:
        errors.append(
            "This file has column(s) AEGIS does not recognise: " + ", ".join(unexpected) + ". "
            "Check for typos against the blank template's headers."
        )
    return errors


# -- row -> entry -------------------------------------------------------------

def _row_is_blank(row_values) -> bool:
    return all(_normalize_cell(v) is None for v in row_values)


def _get(row_values, index):
    if index is None or index >= len(row_values):
        return None
    return row_values[index]


def _entry_from_row(row_values, header_index: dict) -> tuple[dict, list[int]]:
    """Returns (entry, child_positions) -- the second element is the real
    1-5 paper-form slot number behind each item in entry["children"], for
    imports/parsing.py's child-slot error labelling. Blank slots are
    dropped entirely (docs/IMPORT_TEMPLATE.md: never padded to five), so a
    filled "Child 2" following a blank "Child 1" ends up at list position 0
    -- child_positions is what lets an error about it still say "Child 2
    Name", not "Child 1 Name".
    """
    entry = {}
    for header, key in _FIELD_COLUMNS:
        entry[key] = _normalize_cell(_get(row_values, header_index[header]))

    children = []
    child_positions = []
    for n in range(1, CHILD_SLOTS + 1):
        name = _normalize_cell(_get(row_values, header_index[f"Child {n} Name"]))
        dob = _normalize_cell(_get(row_values, header_index[f"Child {n} Date of Birth"]))
        if name is None and dob is None:
            continue
        children.append({"full_name": name, "date_of_birth": dob})
        child_positions.append(n)
    entry["children"] = children

    committees = []
    grievance_ticked = False
    for name in ALL_COMMITTEE_NAMES:
        if _normalize_cell(_get(row_values, header_index[name])) is not None:
            committees.append(name)
            if name == APPOINTED_ONLY_COMMITTEE_NAME:
                grievance_ticked = True
    entry["committees"] = committees

    uncertain_fields = _parse_flagged_fields(
        _normalize_cell(_get(row_values, header_index["Flagged Fields"]))
    )
    if grievance_ticked:
        # "worth catching": Grievance and Reconciliation is Board-appointed
        # and is never printed on the profiling form at all, so a tick here
        # is either a stray mark or a genuine sixth committee the reviewer
        # needs to see -- either way, AEGIS itself can and should flag this
        # (unlike a free-text committees list, a dedicated column means we
        # do not have to rely on the AI remembering to say so).
        uncertain_fields.append({
            "field": "committees",
            "why": (
                "'Grievance and Reconciliation' is ticked on this row, but it is "
                "Board-appointed and is not printed on the profiling form -- confirm this "
                "was actually marked on the paper before approving."
            ),
        })
    entry["uncertain_fields"] = uncertain_fields

    return entry, child_positions


def _label_for_row(entry, row_number: int) -> str:
    if entry.get("source_image"):
        return f"row {row_number} ({entry['source_image']})"
    return f"row {row_number}"


def _rows_to_entries(all_rows: list[list]) -> list[dict]:
    """`all_rows[0]` is the header row; every row after it is one form.
    Returns validated entries or raises ImportValidationError naming the
    spreadsheet row (2-based: row 1 is the header, so the first form is row
    2) and the column header -- not a JSON key, not a zero-based index.
    """
    if not all_rows:
        raise ImportValidationError(["The file has no header row."])

    header_row = all_rows[0]
    header_index: dict[str, int] = {}
    for index, cell in enumerate(header_row):
        text = _normalize_cell(cell)
        if text and text not in header_index:
            header_index[text] = index

    header_errors = _validate_headers(set(header_index))
    if header_errors:
        raise ImportValidationError(header_errors)

    entries = []
    row_numbers = []
    child_positions_by_entry = []
    for offset, row_values in enumerate(all_rows[1:], start=2):
        if _row_is_blank(row_values):
            continue
        entry, child_positions = _entry_from_row(row_values, header_index)
        entries.append(entry)
        row_numbers.append(offset)
        child_positions_by_entry.append(child_positions)

    if not entries:
        raise ImportValidationError(["The file contains no entries."])

    errors: list[str] = []
    for entry, row_number, child_positions in zip(entries, row_numbers, child_positions_by_entry):
        _validate_entry(
            entry, _label_for_row(entry, row_number), errors,
            field_labels=FIELD_LABELS, child_positions=child_positions,
        )

    if errors:
        raise ImportValidationError(errors)

    return entries


# -- format sniffing + top-level entry points --------------------------------

def sniff_import_format(raw_bytes: bytes) -> str:
    """'xlsx', 'csv', or 'json', decided from the bytes themselves -- never
    the filename or a declared content-type, which a volunteer's save-as
    dialog or an email attachment can get wrong without them noticing.
    .xlsx files are zip archives (a bare zip magic number is enough to tell
    them from a CSV, which cannot start with those four bytes and stay
    readable as delimited text); JSON is recognised by its own required
    first non-whitespace character; anything else is treated as CSV, the
    plainest format, so it gets a clear "wrong columns" message instead of
    a format-detection error that means nothing to a volunteer.
    """
    if raw_bytes[:4] == b"PK\x03\x04":
        return "xlsx"
    stripped = raw_bytes.lstrip()
    if stripped[:1] in (b"[", b"{"):
        return "json"
    return "csv"


def _parse_xlsx(raw_bytes: bytes) -> list[dict]:
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ImportValidationError(
            [f"This .xlsx file could not be read ({exc}). Re-save it from Excel and try again."]
        ) from exc
    sheet = workbook.active
    all_rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    return _rows_to_entries(all_rows)


def _parse_csv(raw_bytes: bytes) -> list[dict]:
    try:
        text = raw_bytes.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ImportValidationError([f"The file is not valid UTF-8 text ({exc})."]) from exc
    all_rows = list(csv.reader(io.StringIO(text)))
    return _rows_to_entries(all_rows)


def parse_import_file(raw_bytes: bytes) -> list[dict]:
    """The one entry point imports/admin.py's upload_view calls. Detects
    .xlsx / .csv / .json from the bytes (sniff_import_format) and routes to
    the matching parser, all of which return the identical entry shape and
    raise the identical ImportValidationError -- the caller does not need
    to know or care which format a volunteer uploaded.
    """
    fmt = sniff_import_format(raw_bytes)
    if fmt == "json":
        return parse_import_json(raw_bytes)
    if fmt == "xlsx":
        return _parse_xlsx(raw_bytes)
    return _parse_csv(raw_bytes)


# -- blank template -----------------------------------------------------------

def build_blank_template_workbook():
    """An .xlsx with just AEGIS's expected header row, so the AI (via CSV,
    see docs/IMPORT_TEMPLATE.md) or a volunteer has the right layout to fill
    in rather than inventing one."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Member Profiling Import"
    sheet.append(list(HEADER_TEXTS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    return workbook


def write_blank_template_bytes() -> bytes:
    buffer = io.BytesIO()
    build_blank_template_workbook().save(buffer)
    return buffer.getvalue()
